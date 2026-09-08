import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from analysis.value_analysis import ValueAnalyzer
from data.fetchers.bist_fetcher import BistFetcher


class ExcelGenerator:
    """Excel rapor uretici."""

    def __init__(self):
        self.va = ValueAnalyzer()
        self.bist = BistFetcher()

    def generate_value_report(self, symbols: list, output_path: str = None) -> str:
        """Deger analizi Excel raporu olusturur."""
        if output_path is None:
            output_path = f"reports/daily/deger_analizi_{datetime.now().strftime('%Y%m%d')}.xlsx"

        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Deger Analizi"

        # Stil
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="2980B9", end_color="2980B9", fill_type="solid")
        green_fill = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
        red_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
        yellow_fill = PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        # Basliklar
        headers = [
            "Hisse", "Sirket", "Fiyat (TL)", "Piyasa Deg (M TL)",
            "Net Borc (M TL)", "Fiili Dolasim (%)", "Fiili Hisse",
            "Hesaplanan Deger", "Oran", "Sinyal", "Kar Beklentisi (%)",
            "F/K", "PD/DD",
        ]

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        # Veriler
        for row_idx, symbol in enumerate(symbols, 2):
            result = self.va.analyze(symbol)
            if not result:
                continue

            row_data = [
                result["symbol"],
                result.get("description", "")[:40],
                result["current_price"],
                result["market_cap_mn"],
                result["net_debt_mn"],
                result["float_rate"],
                result["floating_shares"],
                result["calculated_value"],
                result["ratio"],
                result["signal"],
                result["expected_return"],
                result.get("pe"),
                result.get("pb"),
            ]

            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")

                # Sinyal rengi
                if col_idx == 10:  # Sinyal sutunu
                    if "SAT" in str(value):
                        cell.fill = red_fill
                    elif "AL" in str(value):
                        cell.fill = green_fill
                    elif "Borclu" in str(value):
                        cell.fill = PatternFill(start_color="D5D8DC", end_color="D5D8DC", fill_type="solid")
                    else:
                        cell.fill = yellow_fill

                # Sayisal format
                if col_idx in [3, 4, 5, 7, 8, 9, 11, 12, 13]:
                    if value is not None:
                        cell.number_format = "#,##0.00"

        # Sutun genislikleri
        widths = [8, 30, 12, 15, 15, 15, 15, 18, 10, 15, 15, 8, 8]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # Sayfa 2 - Sinyal Tablosu
        ws2 = wb.create_sheet("Sinyal Tablosu")
        ws2.cell(row=1, column=1, value="Oran Araligi").font = header_font
        ws2.cell(row=1, column=1).fill = header_fill
        ws2.cell(row=1, column=2, value="Sinyal").font = header_font
        ws2.cell(row=1, column=2).fill = header_fill
        ws2.cell(row=1, column=3, value="Aciklama").font = header_font
        ws2.cell(row=1, column=3).fill = header_fill

        signals = [
            ("> 3,00", "Hemen SAT", "Kesinlikle satilmali"),
            ("> 1,50", "SAT", "Satismci olunmali"),
            ("> 0,99", "Kafana Gore", "Size kalmis"),
            ("> 0,70", "AL", "Alim firsati"),
            ("> 0,000001", "Evi Barki Sat", "Cok buyuk firsat"),
            ("< 0", "Borclu", "Dikkatli ol"),
        ]

        for i, (rng, sig, desc) in enumerate(signals, 2):
            ws2.cell(row=i, column=1, value=rng).border = thin_border
            ws2.cell(row=i, column=2, value=sig).border = thin_border
            ws2.cell(row=i, column=3, value=desc).border = thin_border

        ws2.column_dimensions["A"].width = 20
        ws2.column_dimensions["B"].width = 15
        ws2.column_dimensions["C"].width = 30

        wb.save(output_path)
        return output_path

    def generate_portfolio_report(self, positions: list, output_path: str = None) -> str:
        """Portfoy Excel raporu olusturur."""
        if output_path is None:
            output_path = f"reports/daily/portfoy_{datetime.now().strftime('%Y%m%d')}.xlsx"

        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Portfoy"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        headers = [
            "Hisse", "Miktar", "Ort. Alis (TL)", "Guncel Fiyat (TL)",
            "Deger (TL)", "Kar/Zarar (TL)", "Kar/Zarar (%)",
        ]

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        toplam_deger = 0
        toplam_kar = 0

        for row_idx, pos in enumerate(positions, 2):
            symbol = pos["stock_code"]
            qty = pos["net_quantity"]
            avg_price = pos["avg_buy_price"]

            current = self.bist.get_price(symbol) or avg_price
            deger = current * qty
            kar = (current - avg_price) * qty
            kar_yuzde = ((current / avg_price) - 1) * 100 if avg_price > 0 else 0

            toplam_deger += deger
            toplam_kar += kar

            row_data = [symbol, qty, avg_price, current, deger, kar, kar_yuzde]
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")
                if col_idx in [3, 4, 5, 6, 7] and value is not None:
                    cell.number_format = "#,##0.00"

        # Ozet satiri
        ozet_row = len(positions) + 2
        ws.cell(row=ozet_row, column=1, value="TOPLAM").font = Font(bold=True)
        ws.cell(row=ozet_row, column=5, value=toplam_deger).font = Font(bold=True)
        ws.cell(row=ozet_row, column=5).number_format = "#,##0.00"
        ws.cell(row=ozet_row, column=6, value=toplam_kar).font = Font(bold=True)
        ws.cell(row=ozet_row, column=6).number_format = "#,##0.00"

        widths = [10, 10, 15, 15, 15, 15, 15]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        wb.save(output_path)
        return output_path

    def generate_comparison_report(self, results: list, output_path: str = None) -> str:
        """Karsilastirma Excel raporu olusturur."""
        if output_path is None:
            output_path = f"reports/daily/karsilastirma_{datetime.now().strftime('%Y%m%d')}.xlsx"

        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Karsilastirma"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="8E44AD", end_color="8E44AD", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        headers = [
            "Hisse", "Fiyat (TL)", "Deger (TL)", "Oran",
            "Sinyal", "Kar (%)", "Piyasa Deg (M)", "F/K", "PD/DD",
        ]

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        for row_idx, r in enumerate(results, 2):
            row_data = [
                r["symbol"], r["current_price"], r["calculated_value"],
                r["ratio"], r["signal"], r["expected_return"],
                r["market_cap_mn"], r.get("pe"), r.get("pb"),
            ]
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")

        wb.save(output_path)
        return output_path


if __name__ == "__main__":
    gen = ExcelGenerator()
    path = gen.generate_value_report(["A1CAP", "THYAO", "GARAN", "AKBNK"])
    print(f"Excel olusturuldu: {path}")
